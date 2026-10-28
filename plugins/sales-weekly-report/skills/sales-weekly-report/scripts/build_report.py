#!/usr/bin/env python3
"""판매관리.xlsx -> "N월 M주차.xlsx" 주간 리포트 생성기.

원본은 절대 수정하지 않는다. 결과는 원본과 같은 폴더에 새 파일로 저장한다.

usage:
  python build_report.py --folder <판매관리.xlsx가 있는 폴더> [--asof 2026-08-30]
                         [--lead-time 7] [--out-dir <저장폴더>] [--source <파일경로>]
"""
import argparse, collections, datetime as dt, glob, json, os, shutil, subprocess, sys, tempfile, unicodedata

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SALES_SHEET = "판매"
STOCK_SHEET = "재고"
SALES_COLS = ["날짜", "품목명", "카테고리", "단가", "판매수량", "매출액", "판매채널"]
STOCK_COLS = ["품목명", "카테고리", "현재재고", "안전재고", "재고상태",
              "최근입고일", "입고수량", "단가", "재고금액", "보관위치"]


def nfc(s):
    """macOS는 한글 파일/시트명을 NFD로 저장한다. 비교 전에 항상 NFC로 정규화한다.
    이걸 빼먹으면 눈으로 같아 보이는 문자열이 != 로 갈린다."""
    return unicodedata.normalize("NFC", s) if isinstance(s, str) else s


class StructureError(Exception):
    """구조가 기대와 다를 때. 추측하지 말고 사용자에게 보고하기 위한 예외."""


# ---------------------------------------------------------------- 입력 로딩
def find_source(folder, explicit=None):
    """macOS 한글 파일명은 NFD 정규화라 리터럴 경로 매칭이 실패한다. glob으로 찾는다."""
    if explicit and os.path.exists(explicit):
        return explicit
    cands = [p for p in glob.glob(os.path.join(folder, "*.xlsx"))
             if not os.path.basename(p).startswith("~$")]
    for p in cands:
        if "판매관리" in nfc(os.path.basename(p)):
            return p
    raise StructureError(
        f"'판매관리'가 들어간 .xlsx를 찾지 못했습니다. 폴더 내 xlsx: "
        f"{[os.path.basename(p) for p in cands] or '없음'}")


def load(src):
    wb = openpyxl.load_workbook(src, data_only=True)
    names = {nfc(n): n for n in wb.sheetnames}
    if SALES_SHEET not in names:
        raise StructureError(f"'{SALES_SHEET}' 시트가 없습니다. 실제 시트: {wb.sheetnames}")

    def grab(name, expected):
        ws = wb[names[name]]
        head = [nfc(c.value) for c in ws[1]][:len(expected)]
        if head != expected:
            raise StructureError(
                f"'{name}' 시트 컬럼이 기대와 다릅니다.\n  기대: {expected}\n  실제: {head}")
        rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
        return rows

    sales = grab(SALES_SHEET, SALES_COLS)
    stock = grab(STOCK_SHEET, STOCK_COLS) if STOCK_SHEET in names else []
    if not sales:
        raise StructureError("'판매' 시트에 데이터 행이 없습니다.")
    return sales, stock


def d(v):
    """날짜 셀이 문자열이든 datetime이든 'YYYY-MM-DD' 문자열로 정규화."""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return str(v)[:10]


# ---------------------------------------------------------------- 주 범위
def week_bounds(asof):
    """주는 월~일. asof가 속한 주의 월요일과 일요일을 돌려준다."""
    mon = asof - dt.timedelta(days=asof.weekday())
    return mon, mon + dt.timedelta(days=6)


def week_label(monday):
    """N주차는 그 주 월요일의 '월 안에서 몇 번째 월요일인가'로 정한다.
    8/3 월 -> 8월 1주차, 8/24 월 -> 8월 4주차. 사용자가 쓰는 관례와 일치한다."""
    return f"{monday.month}월 {(monday.day - 1) // 7 + 1}주차"


# ---------------------------------------------------------------- 리포트 작성
def build(src, out_dir, asof, lead_time):
    sales, stock = load(src)
    for i, r in enumerate(sales):
        sales[i] = (d(r[0]),) + tuple(r[1:])

    mon, sun = week_bounds(asof)
    S, E = mon.strftime("%Y-%m-%d"), sun.strftime("%Y-%m-%d")
    pS = (mon - dt.timedelta(days=7)).strftime("%Y-%m-%d")
    pE = (mon - dt.timedelta(days=1)).strftime("%Y-%m-%d")

    wk = sorted([r for r in sales if S <= r[0] <= E], key=lambda r: (r[0], r[1]))
    prev = [r for r in sales if pS <= r[0] <= pE]
    if not wk:
        alld = sorted({r[0] for r in sales})
        raise StructureError(
            f"대상 주({S}~{E})에 판매 데이터가 없습니다. "
            f"데이터 범위: {alld[0]} ~ {alld[-1]}. --asof 로 기준일을 지정하세요.")

    # 영업일 = 실제로 거래 기록이 있는 날. 휴무일을 분모에 넣으면 일평균이 왜곡된다.
    days = sorted({r[0] for r in wk})
    BD = len(days)

    iq = collections.defaultdict(int)
    ia = collections.defaultdict(int)
    for r in wk:
        iq[r[1]] += r[4]
        ia[r[1]] += r[5]
    items = sorted(iq, key=lambda k: -ia[k])
    chs = sorted({r[6] for r in wk}, key=lambda k: -sum(x[5] for x in wk if x[6] == k))

    short = [r for r in stock if isinstance(r[2], (int, float))
             and isinstance(r[3], (int, float)) and r[2] < r[3]]

    def cover(r):
        """재고 소진 예상일. 판매가 없으면 후순위로 밀되 부족량이 큰 쪽을 먼저."""
        per_day = iq.get(r[0], 0) / BD
        return (r[2] / per_day) if per_day > 0 else float("inf")

    short.sort(key=lambda r: (cover(r), -(r[3] - r[2])))

    wb = Workbook()
    TH = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    H = Font(name="Arial", bold=True, size=13)
    B = Font(name="Arial", bold=True, size=10)
    N = Font(name="Arial", size=10)
    SM = Font(name="Arial", size=9, italic=True, color="666666")
    HF = PatternFill("solid", fgColor="1F4E79")
    SEC = PatternFill("solid", fgColor="DDEBF7")
    RED = PatternFill("solid", fgColor="FFC7CE")
    YEL = PatternFill("solid", fgColor="FFF2CC")
    t = Side(style="thin", color="BFBFBF")
    BD_ = Border(left=t, right=t, top=t, bottom=t)
    C, W, P = "#,##0", "₩#,##0", "0.0%"

    d1 = wb.create_sheet("데이터_이번주판매")
    d1.append(SALES_COLS)
    for r in wk:
        d1.append(list(r))
    d2 = wb.create_sheet("데이터_재고")
    d2.append(STOCK_COLS)
    for r in stock:
        d2.append(list(r))
    for ds in (d1, d2):
        for c in ds[1]:
            c.font, c.fill = TH, HF
            c.alignment = Alignment(horizontal="center")
        ds.freeze_panes = "A2"
        for i in range(1, ds.max_column + 1):
            ds.column_dimensions[get_column_letter(i)].width = 16
    d1.column_dimensions["B"].width = 24
    d2.column_dimensions["A"].width = 24

    L1, L2 = d1.max_row, d2.max_row
    SQ = f"데이터_이번주판매!$E$2:$E${L1}"
    SA = f"데이터_이번주판매!$F$2:$F${L1}"
    SI = f"데이터_이번주판매!$B$2:$B${L1}"
    SC = f"데이터_이번주판매!$G$2:$G${L1}"
    IC, IS_ = f"데이터_재고!$C$2:$C${L2}", f"데이터_재고!$D$2:$D${L2}"
    CNT = f"COUNTA(데이터_이번주판매!$A$2:$A${L1})"

    ws = wb["Sheet"]
    ws.title = "주간리포트"
    ws.column_dimensions["A"].width = 4
    for col, w in zip("BCDEFGHI", [26, 16, 16, 16, 16, 16, 15, 20]):
        ws.column_dimensions[col].width = w
    st = {"r": 2}

    def cell(row, col, val=None):
        return ws.cell(row, col, val)

    def sec(title):
        st["r"] += 1
        cell(st["r"], 2, title).font = H
        for c in range(2, 8):
            cell(st["r"], c).fill = SEC
        st["r"] += 1

    def hdr(cols):
        for i, h in enumerate(cols):
            c = cell(st["r"], 2 + i, h)
            c.font, c.fill = TH, HF
            c.alignment = Alignment(horizontal="center")
            c.border = BD_
        st["r"] += 1

    def note(txt):
        cell(st["r"], 2, txt).font = SM
        st["r"] += 1

    label = week_label(mon)
    cell(st["r"], 2, f"주간 판매 리포트 — {mon.year}년 {label}").font = Font(
        name="Arial", bold=True, size=16)
    st["r"] += 1
    cell(st["r"], 2, f"대상 기간: {S}(월) ~ {E}(일)  ·  실판매 데이터: {days[0]} ~ "
                     f"{days[-1]} ({BD}영업일)").font = N
    st["r"] += 1
    cell(st["r"], 2, f"출처: {os.path.basename(src)} (원본 무수정) · 주 기준 월~일 · "
                     f"재고는 원본 파일 기준 스냅샷").font = SM
    st["r"] += 2

    # --- 1. 요약
    sec("1. 요약")
    hdr(["지표", "값", "비고"])
    first = st["r"]
    prev_total = sum(x[5] for x in prev)
    summary = [
        ("이번 주 총 매출액", f"=SUM({SA})", W, f"{days[0]}~{days[-1]} 합계"),
        ("이번 주 총 판매 건수", f"={CNT}", C, "거래 행 수"),
        ("이번 주 총 판매 수량", f"=SUM({SQ})", C, "개"),
        ("일평균 매출액", f"=SUM({SA})/{BD}", W, f"{BD}영업일 기준"),
        ("건당 평균 매출액", f"=IFERROR(SUM({SA})/{CNT},0)", W, ""),
        ("재고 부족 품목 수", f"=SUMPRODUCT(--({IC}<{IS_}))" if stock else 0, C,
         "현재재고 < 안전재고"),
        (f"전주({pS}~{pE}) 매출액", prev_total, W, "비교용 · 원본 데이터 집계"),
        ("전주 대비 증감률", None, P, ""),
    ]
    for lab, val, fmt, memo in summary:
        cell(st["r"], 2, lab).font = B
        cell(st["r"], 2).border = BD_
        if lab == "전주 대비 증감률":
            val = f'=IFERROR(C{first}/C{st["r"]-1}-1,"")'
        c = cell(st["r"], 3, val)
        c.font, c.number_format, c.border = N, fmt, BD_
        c.alignment = Alignment(horizontal="right")
        cell(st["r"], 4, memo).font = SM
        cell(st["r"], 4).border = BD_
        if lab in ("이번 주 총 매출액", "재고 부족 품목 수"):
            cell(st["r"], 2).fill = YEL
            cell(st["r"], 3).fill = YEL
        st["r"] += 1
    st["r"] += 1

    # --- 2. 판매 분석
    sec("2. 판매 분석")
    cell(st["r"], 2, "2-1. 품목별 판매 실적 (매출액 내림차순)").font = B
    st["r"] += 1
    hdr(["품목명", "판매수량", "매출액", "매출 비중", "건수"])
    a = st["r"]
    for k in items:
        r = st["r"]
        cell(r, 2, k).font = N
        cell(r, 2).border = BD_
        cell(r, 3, f"=SUMIF({SI},$B{r},{SQ})").number_format = C
        cell(r, 4, f"=SUMIF({SI},$B{r},{SA})").number_format = W
        cell(r, 5, f"=IFERROR(D{r}/SUM({SA}),0)").number_format = P
        cell(r, 6, f"=COUNTIF({SI},$B{r})").number_format = C
        for cc in range(3, 7):
            cell(r, cc).font = N
            cell(r, cc).border = BD_
        st["r"] += 1
    b = st["r"] - 1
    cell(st["r"], 2, "합계").font = B
    cell(st["r"], 2).border = BD_
    for col, fm in ((3, C), (4, W), (5, P), (6, C)):
        L = get_column_letter(col)
        c = cell(st["r"], col, f"=SUM({L}{a}:{L}{b})")
        c.font, c.number_format, c.border, c.fill = B, fm, BD_, SEC
    st["r"] += 2

    cell(st["r"], 2, "2-2. 판매채널별 비중 (매출액 내림차순)").font = B
    st["r"] += 1
    hdr(["판매채널", "건수", "판매수량", "매출액", "매출 비중"])
    a = st["r"]
    for k in chs:
        r = st["r"]
        cell(r, 2, k).font = N
        cell(r, 2).border = BD_
        cell(r, 3, f"=COUNTIF({SC},$B{r})").number_format = C
        cell(r, 4, f"=SUMIF({SC},$B{r},{SQ})").number_format = C
        cell(r, 5, f"=SUMIF({SC},$B{r},{SA})").number_format = W
        cell(r, 6, f"=IFERROR(E{r}/SUM({SA}),0)").number_format = P
        for cc in range(3, 7):
            cell(r, cc).font = N
            cell(r, cc).border = BD_
        st["r"] += 1
    b = st["r"] - 1
    cell(st["r"], 2, "합계").font = B
    cell(st["r"], 2).border = BD_
    for col, fm in ((3, C), (4, C), (5, W), (6, P)):
        L = get_column_letter(col)
        c = cell(st["r"], col, f"=SUM({L}{a}:{L}{b})")
        c.font, c.number_format, c.border, c.fill = B, fm, BD_, SEC
    st["r"] += 2

    # --- 3. 재고 현황
    sec("3. 재고 현황 — 안전재고 미달 품목")
    names = [x[0] for x in stock]
    if not short:
        note("안전재고 미달 품목이 없습니다. (현재재고 ≥ 안전재고, 전 품목)")
        st["r"] += 1
    else:
        note(f"긴급 순서 = 재고 소진 예상일(현재재고 ÷ 이번 주 일평균 판매량, "
             f"{BD}영업일 기준) 오름차순. 이번 주 판매 실적이 없는 품목은 \"-\"로 표시하고 후순위.")
        hdr(["품목명", "현재재고", "안전재고", "부족수량", "주간판매량",
             "일평균판매", "소진예상(일)", "권장발주량"])
        a = st["r"]
        for it in short:
            r = st["r"]
            ir = names.index(it[0]) + 2
            cell(r, 2, it[0]).font = B
            cell(r, 2).border = BD_
            cell(r, 2).fill = RED
            cell(r, 3, f"=데이터_재고!C{ir}").number_format = C
            cell(r, 4, f"=데이터_재고!D{ir}").number_format = C
            cell(r, 5, f"=D{r}-C{r}").number_format = C
            cell(r, 6, f"=SUMIF({SI},$B{r},{SQ})").number_format = C
            cell(r, 7, f"=F{r}/{BD}").number_format = "0.0"
            cell(r, 8, f'=IF(G{r}=0,"-",C{r}/G{r})').number_format = "0.0"
            cell(r, 9, f"=E{r}+ROUND(G{r}*{lead_time},0)").number_format = C
            for cc in range(3, 10):
                cell(r, cc).font = N
                cell(r, cc).border = BD_
            st["r"] += 1
        b = st["r"] - 1
        cell(st["r"], 2, "부족 품목 수 / 총 부족수량").font = B
        cell(st["r"], 2).border = BD_
        for col, f_ in ((3, f"=COUNTA(B{a}:B{b})"), (5, f"=SUM(E{a}:E{b})")):
            c = cell(st["r"], col, f_)
            c.font, c.number_format, c.border = B, C, BD_
        st["r"] += 1
        note(f"권장발주량 = 부족수량 + (일평균판매 × {lead_time}일). "
             f"리드타임 {lead_time}일 가정 — 실제 조달 리드타임에 맞춰 조정 필요.")
        st["r"] += 1

    # --- 4. 액션 아이템
    sec("4. 액션 아이템")
    hr = st["r"]
    hdr(["우선순위", "조치 사항", "", "", "근거", "", "", "담당/기한(입력)"])
    ws.merge_cells(start_row=hr, start_column=3, end_row=hr, end_column=5)
    ws.merge_cells(start_row=hr, start_column=6, end_row=hr, end_column=8)

    acts = []
    if short:
        t0 = short[0]
        acts.append(("P1 (긴급)", f"{t0[0]} 즉시 발주 — 소진 임박",
                     f"현재재고 {t0[2]}개 / 안전재고 {t0[3]}개, 이번 주 {iq.get(t0[0],0)}개 판매"))
        acts.append(("P1 (긴급)", f"안전재고 미달 {len(short)}개 품목 일괄 발주 요청",
                     "3번 표의 권장발주량 참조 · 총 부족수량 합계 기준"))
        stale = sorted(short, key=lambda r: d(r[5]))[0]
        acts.append(("P2 (금주)", f"장기 미입고 품목 입고일 점검 ({stale[0]} 등)",
                     f"최근입고일 {d(stale[5])} 이후 재입고 없음"))
    else:
        acts.append(("P2 (금주)", "안전재고 기준선 재검토",
                     "미달 품목 0건 — 기준이 낮게 잡혀 있는지 확인"))
    acts.append(("P2 (금주)", f"매출 상위 품목({', '.join(items[:3])}) 재고 사전 확보",
                 "이번 주 매출 상위 3개 품목 — 결품 시 손실 가장 큼"))
    if chs:
        acts.append(("P3 (차주)", f"{chs[-1]} 채널 부진 원인 점검 및 개선안 수립",
                     f"{chs[-1]} 매출 비중이 {len(chs)}개 채널 중 최하위"))
    acts.append(("P3 (차주)", f"매출 하위 품목({', '.join(items[-3:])}) 재고·프로모션 재검토",
                 "회전율 낮음 — 과잉재고 리스크"))
    wk_total = sum(x[5] for x in wk)
    delta = (wk_total / prev_total - 1) if prev_total else None
    acts.append(("P3 (차주)", "전주 대비 매출 변동 원인 리뷰",
                 f"전주 대비 {delta*100:+.1f}%" if delta is not None
                 else "전주 데이터 없음 — 비교 불가"))

    for p, a_, why in acts:
        r = st["r"]
        cell(r, 2, p).font = B
        cell(r, 2).border = BD_
        cell(r, 2).fill = RED if p.startswith("P1") else (YEL if p.startswith("P2") else SEC)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        cell(r, 3, a_).font = N
        cell(r, 3).alignment = Alignment(vertical="center", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
        cell(r, 6, why).font = SM
        cell(r, 6).alignment = Alignment(vertical="center", wrap_text=True)
        for cc in range(3, 10):
            cell(r, cc).border = BD_
        ws.row_dimensions[r].height = 30
        st["r"] += 1
    st["r"] += 1
    cell(st["r"], 2, f"※ 가정: 주 기준 월~일. 거래 기록이 있는 날만 영업일로 집계"
                     f"({BD}일). 리드타임 {lead_time}일 가정. 재고는 원본 파일의 스냅샷 값.").font = SM

    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False

    out = os.path.join(out_dir, f"{label}.xlsx")
    wb.save(out)
    return out, {
        "week_label": label, "period": f"{S}~{E}",
        "data_range": f"{days[0]}~{days[-1]}", "business_days": BD,
        "total_sales": wk_total, "tx_count": len(wk),
        "total_qty": sum(x[4] for x in wk),
        "prev_week_sales": prev_total,
        "wow_pct": round(delta * 100, 1) if delta is not None else None,
        "shortage_count": len(short),
        "shortage_items": [x[0] for x in short],
        "top_items": [(k, ia[k]) for k in items[:3]],
        "channels": [(k, sum(x[5] for x in wk if x[6] == k)) for k in chs],
    }


# ---------------------------------------------------------------- 재계산
def recalc(path, script, timeout=240):
    """LibreOffice 재계산. 마운트된 사용자 폴더에서 직접 돌리면 타임아웃이 나므로
    로컬 임시 폴더로 복사해서 계산한 뒤 되돌려 놓는다."""
    if not os.path.exists(script):
        return {"status": "skipped", "reason": f"recalc.py 없음: {script}"}
    tmpdir = tempfile.mkdtemp()
    tmp = os.path.join(tmpdir, "w.xlsx")
    shutil.copy(path, tmp)
    try:
        p = subprocess.run([sys.executable, script, tmp, str(timeout)],
                           capture_output=True, text=True, timeout=timeout + 60)
        try:
            res = json.loads(p.stdout)
        except Exception:
            return {"status": "unknown", "stdout": p.stdout[-500:], "stderr": p.stderr[-500:]}
        if res.get("status") == "success":
            shutil.copy(tmp, path)
        return res
    except subprocess.TimeoutExpired:
        return {"status": "timeout"}
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--folder", required=True)
    ap.add_argument("--source")
    ap.add_argument("--out-dir")
    ap.add_argument("--asof", help="기준일 YYYY-MM-DD (기본: 오늘)")
    ap.add_argument("--lead-time", type=int, default=7)
    ap.add_argument("--recalc-script",
                    default="/sessions/*/mnt/.claude/skills/xlsx/scripts/recalc.py")
    a = ap.parse_args()
    try:
        src = find_source(a.folder, a.source)
        asof = (dt.datetime.strptime(a.asof, "%Y-%m-%d").date() if a.asof
                else dt.date.today())
        out_dir = a.out_dir or os.path.dirname(src)
        out, stats = build(src, out_dir, asof, a.lead_time)
    except StructureError as e:
        print(json.dumps({"status": "structure_error", "message": str(e)},
                         ensure_ascii=False, indent=2))
        sys.exit(2)
    rs = a.recalc_script
    if "*" in rs:
        hits = glob.glob(rs)
        rs = hits[0] if hits else rs
    stats["source"] = src
    stats["output"] = out
    stats["recalc"] = recalc(out, rs)
    print(json.dumps(stats, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
